import base64
import io
import re
import zipfile
from pathlib import Path
from PIL import Image, ImageOps
from .errors import AppError

MAX_BYTES=8*1024*1024
MAX_TEXT=16000
Image.MAX_IMAGE_PIXELS=20_000_000

def parse_file(name,data):
    if not data or len(data)>MAX_BYTES: raise AppError('attachment_too_large','Размер файла должен быть от 1 байта до 8 МБ.',413)
    name=Path(name.replace('\\','/')).name[:160]
    suffix=Path(name).suffix.lower(); warnings=[]; image=None
    try:
        if suffix in {'.docx','.xlsx'}:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist())>32*1024*1024:
                    raise AppError('attachment_too_large','Слишком большой распакованный документ.',413)
        if suffix in {'.jpg','.jpeg','.png'}:
            with Image.open(io.BytesIO(data)) as img:
                if img.format not in {'JPEG','PNG'}: raise ValueError('image signature')
                img=ImageOps.exif_transpose(img).convert('RGB'); img.thumbnail((1600,1600))
                out=io.BytesIO(); img.save(out,format='JPEG',quality=85)
                image='data:image/jpeg;base64,'+base64.b64encode(out.getvalue()).decode('ascii')
            text='Изображение будет распознано моделью при отправке сообщения.'
        elif suffix=='.pdf':
            from pypdf import PdfReader
            if not data.startswith(b'%PDF'): raise ValueError('PDF signature')
            reader=PdfReader(io.BytesIO(data))
            if reader.is_encrypted: raise AppError('invalid_attachment','Снимите пароль с PDF перед загрузкой.',422)
            if len(reader.pages)>30: raise AppError('attachment_too_large','Максимум 30 страниц PDF.',413)
            text='\n'.join(page.extract_text() or '' for page in reader.pages)
            if not text.strip():
                raise AppError('invalid_attachment','PDF состоит из сканов. Прикрепите страницу как JPEG/PNG для распознавания.',422)
        elif suffix=='.docx':
            from docx import Document
            doc=Document(io.BytesIO(data))
            text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        elif suffix=='.xlsx':
            import openpyxl
            book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
            rows=[]
            for sheet in book.worksheets[:5]:
                rows.append(sheet.title)
                for row in sheet.iter_rows(max_row=500,max_col=20,values_only=True):
                    rows.append(' | '.join(str(c) for c in row if c is not None))
            book.close(); text='\n'.join(rows)
            warnings.append('Для поиска прочитаны первые 5 листов, до 500 строк и 20 столбцов каждого.')
        elif suffix=='.xls':
            import xlrd
            book=xlrd.open_workbook(file_contents=data,on_demand=True)
            text='\n'.join(' | '.join(str(c) for c in sheet.row_values(r)[:20]) for sheet in book.sheets()[:5] for r in range(min(sheet.nrows,500)))
            book.release_resources(); warnings.append('Для поиска прочитаны первые 5 листов, до 500 строк и 20 столбцов каждого.')
        else:
            raise AppError('invalid_attachment','Поддерживаются JPEG, PNG, PDF, DOCX, XLSX и XLS. Старый Word DOC сохраните как DOCX.',422)
    except AppError: raise
    except Exception:
        raise AppError('invalid_attachment','Не удалось прочитать файл. Проверьте формат, целостность и отсутствие пароля.',422)
    if len(text)>MAX_TEXT: warnings.append('Текст ограничен первыми 16 000 символами; разделите длинную спецификацию.')
    text=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',text[:MAX_TEXT])
    return {'name':name,'extracted_text':text,'kind':'image' if image else 'document','warnings':warnings,'image':image}
